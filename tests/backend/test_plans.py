"""Plans: Сынақ once per phone, Ұстаз by the month, and counting what is made.

Marked `real_plans`, so the limits apply as in production. Time is moved by
passing `now` to the plan service directly - the rules that depend on the
calendar (periods, carry-over, the summer pause, school-year ends) are tested
there; the HTTP tests check that the routes actually charge and refuse.
"""

import asyncio
import io
import uuid
from datetime import datetime, timedelta, timezone

import pytest
from bson import ObjectId

pytestmark = pytest.mark.real_plans

SPEC = {"subject": "physics", "grade": 8, "topic": "Ньютон заңдары",
        "type": "sor", "difficulty": "medium", "lang": "kk", "count": 2}


def run_db(coro):
    return asyncio.run(coro)


def utc(*args):
    return datetime(*args, tzinfo=timezone.utc)


@pytest.fixture
def plans(backend):
    from app.services import plans as module

    return module


@pytest.fixture
def account(backend):
    from fastapi.testclient import TestClient

    opened = []

    def _make(role="teacher", phone=True):
        client = TestClient(backend.app)
        client.__enter__()
        opened.append(client)
        response = client.post("/api/auth/signup", json={
            "email": f"p{uuid.uuid4().hex[:10]}@school.kz", "password": "password123",
            "name": "Тест", "role": role})
        assert response.status_code == 201, response.text
        user = response.json()["user"]
        if phone:
            number = "7707" + f"{uuid.uuid4().int % 10_000_000:07d}"
            run_db(backend.db.db.users.update_one(
                {"_id": ObjectId(user["id"])},
                {"$set": {"phone": number, "phone_verified_at": datetime.now(timezone.utc)}}))
            user["phone"] = number
        client.user = user
        return client

    yield _make
    for client in opened:
        client.__exit__(None, None, None)


def doc(backend, user_id):
    return run_db(backend.db.db.users.find_one({"_id": ObjectId(user_id)}))


@pytest.fixture(autouse=True)
def _agent_writes_papers(backend, monkeypatch):
    from app.api.assessments import assessment_limiter
    from app.ws.manager import agent_manager

    assessment_limiter._attempts.clear()

    async def paper(spec, timeout):
        return {"questions": [{"text": "Сұрақ бірінші", "marks": 2},
                              {"text": "Сұрақ екінші", "marks": 3}], "total_marks": 5}

    monkeypatch.setattr(agent_manager, "request_assessment", paper)


# ------------------------------------------------------------------ trial --


class TestTrial:
    def test_a_teacher_without_a_phone_is_on_nothing(self, account):
        client = account(phone=False)
        plan = client.get("/api/plan").json()["plan"]
        assert plan["kind"] == "free"
        assert plan["trial_unavailable"] == "phone_required"
        response = client.post("/api/assessments", json=SPEC)
        assert response.status_code == 403 and response.json()["code"] == "plan_required"

    def test_a_confirmed_phone_starts_the_trial(self, account):
        client = account()
        plan = client.get("/api/plan").json()["plan"]
        assert plan["kind"] == "trial"
        assert plan["available"] == {"videos": 2, "assessments": 2, "quiz_drafts": 3, "lesson_plans": 2}
        assert plan["watermark"] is True and plan["docx_signature"] is True

    def test_two_papers_then_no_more(self, account):
        client = account()
        assert client.post("/api/assessments", json=SPEC).status_code == 201
        assert client.post("/api/assessments", json=SPEC).status_code == 201
        third = client.post("/api/assessments", json=SPEC)
        assert third.status_code == 403
        assert third.json()["code"] == "plan_exhausted" and third.json()["resource"] == "assessments"
        assert client.get("/api/plan").json()["plan"]["available"]["assessments"] == 0

    def test_a_failed_paper_is_not_charged(self, account, monkeypatch):
        from app.ws.manager import agent_manager

        async def broken(spec, timeout):
            return {"error": "llm_failed"}

        monkeypatch.setattr(agent_manager, "request_assessment", broken)
        client = account()
        assert client.post("/api/assessments", json=SPEC).status_code >= 400
        assert client.get("/api/plan").json()["plan"]["used"]["assessments"] == 0

    def test_one_trial_per_number_even_after_deleting_the_account(self, backend, account):
        first = account()
        first.get("/api/plan")  # starts the trial, claims the number
        phone = first.user["phone"]
        first.delete("/api/auth/account")

        second = account(phone=False)
        run_db(backend.db.db.users.update_one(
            {"_id": ObjectId(second.user["id"])},
            {"$set": {"phone": phone, "phone_verified_at": datetime.now(timezone.utc)}}))
        plan = second.get("/api/plan").json()["plan"]
        assert plan["kind"] == "free" and plan["trial_unavailable"] == "phone_used"

    def test_only_the_hmac_of_the_number_is_kept(self, backend, account):
        client = account()
        client.get("/api/plan")
        claim = run_db(backend.db.db.trial_claims.find_one({"user_id": client.user["id"]}))
        assert claim and client.user["phone"] not in str(claim)

    def test_students_never_get_a_trial(self, account):
        client = account(role="student")
        plan = client.get("/api/plan").json()["plan"]
        assert plan["kind"] == "free" and plan["trial_unavailable"] == "teacher_only"

    def test_the_trial_caps_quality_at_720p(self, plans):
        trial = run_db(plans.get_plan(plans.TRIAL))
        assert plans.clamp_effort("high", trial) == "medium"
        assert plans.clamp_effort("low", trial) == "low"


# ------------------------------------------------------------- quiz drafts --


class TestQuizDrafts:
    @pytest.fixture(autouse=True)
    def _drafts(self, backend, monkeypatch):
        from app.api.quiz_drafts import draft_limiter
        from app.ws.manager import agent_manager

        draft_limiter._attempts.clear()

        async def write(spec, timeout):
            return {"questions": [{"type": "choice", "text": "Q", "options": ["a", "b"], "answer": 0}],
                    "requested": spec["count"]}

        monkeypatch.setattr(agent_manager, "request_quiz", write)

    BODY = {"subject": "physics", "grade": 8, "topic": "Күш", "lang": "kk", "count": 3}

    def test_three_new_quizzes_on_the_trial(self, account):
        client = account()
        codes = [client.post("/api/quiz-drafts", json=self.BODY).status_code for _ in range(4)]
        assert codes == [200, 200, 200, 403]

    def test_editing_a_quiz_is_not_a_new_one(self, account):
        client = account()
        for _ in range(5):
            response = client.post("/api/quiz-drafts", json={**self.BODY, "count": 1, "purpose": "edit"})
            assert response.status_code == 200
        assert client.get("/api/plan").json()["plan"]["used"]["quiz_drafts"] == 0

    def test_an_edit_still_needs_a_plan_and_is_small(self, account):
        free = account(phone=False)
        assert free.post("/api/quiz-drafts", json={**self.BODY, "purpose": "edit"}).status_code == 403
        client = account()
        assert client.post("/api/quiz-drafts", json={**self.BODY, "count": 12, "purpose": "edit"}).status_code == 400


# ------------------------------------------------------------- the video --


class TestVideos:
    def _ask(self, client, agent_secret, reply):
        with client.websocket_connect("/ws/agent") as agent:
            agent.send_json({"type": "auth", "token": agent_secret})
            assert agent.receive_json().get("type") == "auth_ok"
            with client.websocket_connect("/ws") as ws:
                ws.send_json({"type": "create_chat", "data": {"title": "t"}})
                chat_id = ws.receive_json()["data"]["id"]
                ws.send_json({"type": "user_message", "data": {
                    "chat_id": chat_id, "prompt": f"сұрақ {uuid.uuid4().hex[:8]}", "effort": "high"}})
                first = ws.receive_json()
                if first.get("type") != "message_received":
                    return first, None
                sent = agent.receive_json()
                agent.send_json({"request_id": sent["request_id"], **reply})
                return ws.receive_json(), sent

    def test_a_trial_video_is_watermarked_capped_and_counted(self, account, agent_secret):
        client = account()
        frame, sent = self._ask(client, agent_secret, {
            "status": "complete", "text": "жауап",
            "video_path": "rendered_1_trial.mp4", "clean_video_path": "rendered_1.mp4"})
        assert sent["watermark"] == "Akron · сынақ"
        assert sent["effort"] == "medium"
        assert frame["data"]["video_url"] == "/media/rendered_1_trial.mp4"
        assert client.get("/api/plan").json()["plan"]["used"]["videos"] == 1

    def test_a_failed_video_gives_the_unit_back(self, account, agent_secret):
        client = account()
        self._ask(client, agent_secret, {"status": "error", "error": "render failed"})
        assert client.get("/api/plan").json()["plan"]["used"]["videos"] == 0

    def test_a_text_only_answer_gives_the_unit_back(self, account, agent_secret):
        client = account()
        self._ask(client, agent_secret, {"status": "complete", "text": "тек мәтін", "video_path": ""})
        assert client.get("/api/plan").json()["plan"]["used"]["videos"] == 0

    def test_without_a_plan_the_socket_refuses(self, account, agent_secret):
        client = account(phone=False)
        frame, sent = self._ask(client, agent_secret, {})
        assert frame["data"]["code"] == "plan_required" and sent is None

    def test_ustaz_has_no_watermark_and_full_quality(self, backend, account, agent_secret, plans):
        client = account()
        run_db(plans.grant(doc(backend, client.user["id"]), plans.USTAZ_MONTH))
        frame, sent = self._ask(client, agent_secret, {
            "status": "complete", "text": "жауап", "video_path": "rendered_1.mp4"})
        assert "watermark" not in sent and sent["effort"] == "high"


# ------------------------------------------------------------ the periods --


class TestUstazPeriods:
    def _user(self, backend, account):
        client = account()
        return client, doc(backend, client.user["id"])

    def test_a_month_starts_with_twenty_videos(self, backend, account, plans):
        _, user = self._user(backend, account)
        user = run_db(plans.grant(user, plans.USTAZ_MONTH, now=utc(2026, 10, 5, 9)))
        ent = run_db(plans.resolve(user, now=utc(2026, 10, 6)))
        assert ent.kind == plans.ON_USTAZ and ent.available("videos") == 20
        assert ent.available("quiz_drafts") is None  # unlimited

    def test_unused_videos_carry_into_the_next_month(self, backend, account, plans):
        _, user = self._user(backend, account)
        user = run_db(plans.grant(user, plans.USTAZ_MONTH, months=3, now=utc(2026, 10, 5, 9)))
        run_db(backend.db.db.users.update_one({"_id": user["_id"]}, {"$set": {"period.used.videos": 15}}))
        user = doc(backend, str(user["_id"]))

        ent = run_db(plans.resolve(user, now=utc(2026, 11, 6)))

        assert ent.carryover == 5
        assert ent.available("videos") == 25

    def test_the_carry_over_is_capped(self, backend, account, plans):
        _, user = self._user(backend, account)
        user = run_db(plans.grant(user, plans.USTAZ_MONTH, months=3, now=utc(2026, 10, 5, 9)))
        run_db(backend.db.db.users.update_one({"_id": user["_id"]}, {"$set": {"period.carryover": 20}}))
        user = doc(backend, str(user["_id"]))
        ent = run_db(plans.resolve(user, now=utc(2026, 11, 6)))
        assert ent.carryover == 20  # 20 + 20 unused, capped at 20

    def test_extra_videos_are_spent_after_the_month(self, backend, account, plans):
        _, user = self._user(backend, account)
        user = run_db(plans.grant(user, plans.USTAZ_MONTH))
        run_db(backend.db.db.users.update_one({"_id": user["_id"]},
                                              {"$set": {"period.used.videos": 20}, "$inc": {"extra_videos": 10}}))
        user = doc(backend, str(user["_id"]))
        run_db(plans.reserve(user, "videos", "r1"))
        assert doc(backend, str(user["_id"]))["extra_videos"] == 9
        run_db(plans.release("r1"))
        assert doc(backend, str(user["_id"]))["extra_videos"] == 10

    def test_expiry_falls_back_to_the_trial(self, backend, account, plans):
        _, user = self._user(backend, account)
        user = run_db(plans.grant(user, plans.USTAZ_MONTH, now=utc(2026, 10, 5)))
        ent = run_db(plans.resolve(user, now=utc(2026, 11, 6)))
        assert ent.kind == plans.ON_TRIAL

    def test_the_school_year_ends_on_the_first_of_june(self, backend, account, plans):
        _, user = self._user(backend, account)
        user = run_db(plans.grant(user, plans.USTAZ_YEAR, now=utc(2026, 10, 5)))
        local_end = user["subscription"]["expires_at"].replace(tzinfo=timezone.utc) + timedelta(hours=5)
        assert (local_end.month, local_end.day, local_end.year) == (6, 1, 2027)

    def test_bought_in_july_it_runs_to_the_next_june(self, plans):
        end = plans.school_year_end(utc(2026, 7, 10)) + timedelta(hours=5)
        assert (end.year, end.month, end.day) == (2027, 6, 1)

    def test_bought_in_late_may_it_runs_to_the_next_june(self, plans):
        """Not a fortnight for the price of a year."""
        late = plans.school_year_end(utc(2027, 5, 20)) + timedelta(hours=5)
        early = plans.school_year_end(utc(2027, 5, 10)) + timedelta(hours=5)
        assert (late.year, late.month, late.day) == (2028, 6, 1)
        assert (early.year, early.month, early.day) == (2027, 6, 1)

    def test_granting_again_extends(self, backend, account, plans):
        _, user = self._user(backend, account)
        user = run_db(plans.grant(user, plans.USTAZ_MONTH, now=utc(2026, 10, 5)))
        user = run_db(plans.grant(user, plans.USTAZ_MONTH, months=2, now=utc(2026, 10, 20)))
        expires = user["subscription"]["expires_at"].replace(tzinfo=timezone.utc)
        assert (expires.year, expires.month, expires.day) == (2027, 1, 5)

    def test_month_to_year_switches_and_keeps_the_month_in_progress(self, backend, account, plans):
        _, user = self._user(backend, account)
        user = run_db(plans.grant(user, plans.USTAZ_MONTH, now=utc(2026, 10, 5)))
        run_db(backend.db.db.users.update_one({"_id": user["_id"]}, {"$set": {"period.used.videos": 7}}))
        user = run_db(plans.grant(doc(backend, str(user["_id"])), plans.USTAZ_YEAR, now=utc(2026, 10, 20)))
        assert user["subscription"]["plan_id"] == plans.USTAZ_YEAR
        local_end = user["subscription"]["expires_at"].replace(tzinfo=timezone.utc) + timedelta(hours=5)
        assert (local_end.year, local_end.month, local_end.day) == (2027, 6, 1)
        # Not a fresh month of videos: what was used this month stays used.
        assert user["period"]["used"]["videos"] == 7

    def test_year_to_month_adds_the_month_after_the_paid_year(self, backend, account, plans):
        _, user = self._user(backend, account)
        user = run_db(plans.grant(user, plans.USTAZ_YEAR, now=utc(2026, 10, 5)))
        year_end = user["subscription"]["expires_at"].replace(tzinfo=timezone.utc)
        user = run_db(plans.grant(user, plans.USTAZ_MONTH, now=utc(2026, 10, 20)))
        assert user["subscription"]["plan_id"] == plans.USTAZ_MONTH
        expires = user["subscription"]["expires_at"].replace(tzinfo=timezone.utc)
        assert expires == plans.add_months(year_end, 1)

    def test_switching_back_and_forth_gives_no_new_videos(self, backend, account, plans):
        _, user = self._user(backend, account)
        user = run_db(plans.grant(user, plans.USTAZ_MONTH, now=utc(2026, 10, 5)))
        run_db(backend.db.db.users.update_one({"_id": user["_id"]}, {"$set": {"period.used.videos": 20}}))
        for plan_id in (plans.USTAZ_YEAR, plans.USTAZ_MONTH, plans.USTAZ_YEAR):
            user = run_db(plans.grant(doc(backend, str(user["_id"])), plan_id, now=utc(2026, 10, 6)))
        assert run_db(plans.resolve(user, now=utc(2026, 10, 7))).available("videos") == 0


class TestSummerPause:
    def test_pause_keeps_the_paid_days_until_september(self, backend, account, plans):
        client = account()
        user = run_db(plans.grant(doc(backend, client.user["id"]), plans.USTAZ_MONTH, now=utc(2026, 6, 10)))
        user = run_db(plans.pause(user, now=utc(2026, 6, 20)))
        sub = user["subscription"]
        assert sub["status"] == "paused"
        remaining = sub["remaining_sec"]
        assert remaining == int((utc(2026, 7, 10) - utc(2026, 6, 20)).total_seconds())

        # While paused there is no subscription to generate with.
        assert run_db(plans.resolve(user, now=utc(2026, 7, 1))).kind != plans.ON_USTAZ

        # 1 September: it comes back on its own, with the days that were left.
        ent = run_db(plans.resolve(user, now=utc(2026, 9, 1, 12)))
        assert ent.kind == plans.ON_USTAZ
        expires = ent.user["subscription"]["expires_at"].replace(tzinfo=timezone.utc)
        assert abs((expires - (utc(2026, 8, 31, 19) + timedelta(seconds=remaining))).total_seconds()) < 2

    def test_no_pause_outside_the_summer(self, backend, account, plans):
        client = account()
        user = run_db(plans.grant(doc(backend, client.user["id"]), plans.USTAZ_MONTH, now=utc(2026, 10, 1)))
        with pytest.raises(plans.PlanRefusal) as refused:
            run_db(plans.pause(user, now=utc(2026, 10, 10)))
        assert refused.value.code == "pause_summer_only"

    def test_the_school_year_plan_does_not_pause(self, backend, account, plans):
        client = account()
        user = run_db(plans.grant(doc(backend, client.user["id"]), plans.USTAZ_YEAR, now=utc(2026, 5, 20)))
        with pytest.raises(plans.PlanRefusal) as refused:
            run_db(plans.pause(user, now=utc(2026, 5, 25)))
        assert refused.value.code == "pause_month_only"


class TestReminders:
    def test_mailed_once_three_days_before_the_end(self, backend, account, plans, monkeypatch):
        from app.services import mailer

        sent = []

        async def fake_send(to, subject, text, html=None):
            # Signing up mails a confirmation link through the same function;
            # only the reminders are counted here.
            if "жазылым" in subject:
                sent.append(to)
            return True

        monkeypatch.setattr(mailer, "send", fake_send)
        client = account()
        user = doc(backend, client.user["id"])
        now = datetime.now(timezone.utc).replace(microsecond=0)
        run_db(plans.grant(user, plans.USTAZ_MONTH, until=now + timedelta(days=2)))

        run_db(plans.send_reminders(now=now))
        run_db(plans.send_reminders(now=now + timedelta(hours=1)))

        assert sent.count(client.user["email"]) == 1


# ------------------------------------------------------------------ admin --


class TestAdminPlanActions:
    def _staff(self, backend, account, role):
        staff = account()
        run_db(backend.db.db.users.update_one({"_id": ObjectId(staff.user["id"])},
                                              {"$set": {"staff_role": role}}))
        return staff

    def test_admin_grants_with_a_reason_and_it_is_logged(self, backend, account):
        admin = self._staff(backend, account, "admin")
        teacher = account()
        response = admin.post(f"/api/admin/users/{teacher.user['id']}/subscription",
                              json={"plan_id": "ustaz_month", "months": 2, "reason": "Мектеп пилоты"})
        assert response.status_code == 200
        assert teacher.get("/api/plan").json()["plan"]["kind"] == "ustaz"
        entry = run_db(backend.db.db.admin_audit.find_one({"action": "plan.grant",
                                                            "target.user_id": teacher.user["id"]}))
        assert entry["reason"] == "Мектеп пилоты" and entry["actor"]["email"] == admin.user["email"]

    def test_no_reason_no_grant(self, backend, account):
        admin = self._staff(backend, account, "admin")
        teacher = account()
        response = admin.post(f"/api/admin/users/{teacher.user['id']}/subscription",
                              json={"plan_id": "ustaz_month", "reason": "  "})
        assert response.status_code == 400 and response.json()["code"] == "reason_required"

    def test_support_cannot_grant_but_can_reset_a_trial(self, backend, account):
        support = self._staff(backend, account, "support")
        teacher = account()
        teacher.post("/api/assessments", json=SPEC)
        assert support.post(f"/api/admin/users/{teacher.user['id']}/subscription",
                            json={"plan_id": "ustaz_month", "reason": "x"}).status_code == 403

        assert support.post(f"/api/admin/users/{teacher.user['id']}/trial/reset",
                            json={"reason": "Қате тіркелді"}).status_code == 200
        plan = teacher.get("/api/plan").json()["plan"]
        assert plan["kind"] == "trial" and plan["used"]["assessments"] == 0

    def test_blocking_signs_out_and_keeps_out(self, backend, account):
        support = self._staff(backend, account, "support")
        teacher = account()
        assert support.post(f"/api/admin/users/{teacher.user['id']}/block",
                            json={"reason": "Пробный пакет фармы"}).status_code == 200
        assert teacher.get("/api/auth/me").status_code == 401
        login = teacher.post("/api/auth/login", json={"email": teacher.user["email"], "password": "password123"})
        assert login.status_code == 403 and login.json()["code"] == "account_blocked"

        support.post(f"/api/admin/users/{teacher.user['id']}/unblock", json={"reason": "Қате"})
        login = teacher.post("/api/auth/login", json={"email": teacher.user["email"], "password": "password123"})
        assert login.status_code == 200

    def test_staff_cannot_be_blocked_from_the_panel(self, backend, account):
        support = self._staff(backend, account, "support")
        admin = self._staff(backend, account, "admin")
        response = support.post(f"/api/admin/users/{admin.user['id']}/block", json={"reason": "x"})
        assert response.status_code == 400

    def test_limits_change_without_a_deploy(self, backend, account):
        admin = self._staff(backend, account, "admin")
        try:
            response = admin.put("/api/admin/plans/trial",
                                 json={"limits": {"assessments": 5}, "reason": "Тест"})
            assert response.status_code == 200
            teacher = account()
            assert teacher.get("/api/plan").json()["plan"]["available"]["assessments"] == 5
        finally:
            run_db(backend.db.db.plans.delete_one({"_id": "trial"}))

    def test_an_unknown_limit_is_refused(self, backend, account):
        admin = self._staff(backend, account, "admin")
        response = admin.put("/api/admin/plans/trial", json={"limits": {"rockets": 1}, "reason": "x"})
        assert response.status_code == 400


# ------------------------------------------------------------------- docx --


class TestDocx:
    def _paper(self, client):
        response = client.post("/api/assessments", json=SPEC)
        assert response.status_code == 201, response.text
        return response.json()["id"]

    def _footer(self, content):
        from docx import Document

        document = Document(io.BytesIO(content))
        footer = " ".join(p.text for s in document.sections for p in s.footer.paragraphs)
        body = "\n".join(p.text for p in document.paragraphs)
        return footer, body

    def test_a_trial_paper_is_signed(self, account):
        client = account()
        response = client.get(f"/api/assessments/{self._paper(client)}/docx")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml")
        footer, body = self._footer(response.content)
        assert "Akron арқылы жасалды" in footer
        assert "БЖБ" in body and "Сұрақ бірінші" in body and "Жалпы: 5 балл" in body

    def test_a_ustaz_paper_is_not(self, backend, account, plans):
        client = account()
        run_db(plans.grant(doc(backend, client.user["id"]), plans.USTAZ_MONTH))
        footer, _ = self._footer(client.get(f"/api/assessments/{self._paper(client)}/docx").content)
        assert "Akron" not in footer

    def test_only_the_owner_downloads_it(self, account):
        owner = account()
        paper = self._paper(owner)
        other = account()
        assert other.get(f"/api/assessments/{paper}/docx").status_code == 404


# ------------------------------------------------- papers, quizzes, limits --


def _staff(backend, account, role="admin"):
    staff = account()
    run_db(backend.db.db.users.update_one({"_id": ObjectId(staff.user["id"])},
                                          {"$set": {"staff_role": role}}))
    return staff


class TestPaperCarryOver:
    """БЖБ/ТЖБ come in a rush at the end of a term, so unused ones carry."""

    def _month(self, backend, account, plans, months=3):
        client = account()
        user = run_db(plans.grant(doc(backend, client.user["id"]), plans.USTAZ_MONTH,
                                  months=months, now=utc(2026, 10, 5, 9)))
        return client, user

    def test_a_month_has_twenty_papers_and_the_unused_carry(self, backend, account, plans):
        _, user = self._month(backend, account, plans)
        assert run_db(plans.resolve(user, now=utc(2026, 10, 6))).available("assessments") == 20
        run_db(backend.db.db.users.update_one({"_id": user["_id"]}, {"$set": {"period.used.assessments": 15}}))

        ent = run_db(plans.resolve(doc(backend, str(user["_id"])), now=utc(2026, 11, 6)))

        assert ent.carry("assessments") == 5
        assert ent.available("assessments") == 25
        assert ent.carryover == 20  # videos carry as they always did

    def test_the_paper_carry_over_is_capped(self, backend, account, plans):
        _, user = self._month(backend, account, plans)
        run_db(backend.db.db.users.update_one({"_id": user["_id"]}, {"$set": {"period.carried": {"assessments": 20}}}))
        ent = run_db(plans.resolve(doc(backend, str(user["_id"])), now=utc(2026, 11, 6)))
        assert ent.carry("assessments") == 20  # 20 + 20 unused, capped at 20

    def test_carried_papers_can_be_spent_and_are_shown(self, backend, account, plans):
        client = account()
        user = run_db(plans.grant(doc(backend, client.user["id"]), plans.USTAZ_MONTH))
        run_db(backend.db.db.users.update_one({"_id": user["_id"]}, {"$set": {
            "period.carried": {"assessments": 5}, "period.used.assessments": 24}}))

        run_db(plans.reserve(doc(backend, client.user["id"]), "assessments", "carry-25"))
        with pytest.raises(plans.PlanRefusal):
            run_db(plans.reserve(doc(backend, client.user["id"]), "assessments", "carry-26"))
        assert client.get("/api/plan").json()["plan"]["carried"] == {"assessments": 5}

    def test_the_pause_keeps_unused_papers_too(self, backend, account, plans):
        client = account()
        user = run_db(plans.grant(doc(backend, client.user["id"]), plans.USTAZ_MONTH, now=utc(2026, 6, 10)))
        run_db(backend.db.db.users.update_one({"_id": user["_id"]}, {"$set": {"period.used.assessments": 12}}))
        user = run_db(plans.pause(doc(backend, client.user["id"]), now=utc(2026, 6, 20)))
        ent = run_db(plans.resolve(user, now=utc(2026, 9, 1, 12)))
        assert ent.kind == plans.ON_USTAZ and ent.carry("assessments") == 8


class TestQuizFairUse:
    """Unlimited on the page; a ceiling no single teacher reaches."""

    @pytest.fixture(autouse=True)
    def _drafts(self, backend, monkeypatch):
        from app.api.quiz_drafts import draft_limiter
        from app.ws.manager import agent_manager

        draft_limiter._attempts.clear()

        async def write(spec, timeout):
            return {"questions": [{"type": "choice", "text": "Q", "options": ["a", "b"], "answer": 0}],
                    "requested": spec["count"]}

        monkeypatch.setattr(agent_manager, "request_quiz", write)

    BODY = {"subject": "physics", "grade": 8, "topic": "Күш", "lang": "kk", "count": 3}

    def test_unlimited_until_the_ceiling(self, backend, account, plans):
        client = account()
        user = run_db(plans.grant(doc(backend, client.user["id"]), plans.USTAZ_MONTH))
        shown = client.get("/api/plan").json()["plan"]
        assert shown["available"]["quiz_drafts"] is None
        assert shown["fair_use"] == {"quiz_drafts": 100}

        run_db(backend.db.db.users.update_one({"_id": user["_id"]}, {"$set": {"period.used.quiz_drafts": 99}}))
        assert client.post("/api/quiz-drafts", json=self.BODY).status_code == 200
        refused = client.post("/api/quiz-drafts", json=self.BODY)

        assert refused.status_code == 403
        assert refused.json()["code"] == "fair_use_reached" and refused.json()["resource"] == "quiz_drafts"
        # Still unlimited on the page: the ceiling is named there, not counted down.
        assert client.get("/api/plan").json()["plan"]["available"]["quiz_drafts"] is None

    def test_the_trial_keeps_its_three(self, account):
        shown = account().get("/api/plan").json()["plan"]
        assert shown["available"]["quiz_drafts"] == 3 and shown["fair_use"] == {}


class TestLimitHits:
    """Whether a limit is the right size is answered by who runs out, and when."""

    def test_running_out_on_the_trial_is_counted(self, backend, account):
        client = account()
        for _ in range(3):
            client.post("/api/assessments", json=SPEC)
        hit = run_db(backend.db.db.plan_limit_hits.find_one({"user_id": client.user["id"]}))
        assert (hit["kind"], hit["resource"], hit["code"]) == ("trial", "assessments", "plan_exhausted")
        assert "period_day" not in hit

    def test_running_out_on_ustaz_says_how_far_into_the_month(self, backend, account, plans):
        client = account()
        user = run_db(plans.grant(doc(backend, client.user["id"]), plans.USTAZ_MONTH))
        run_db(backend.db.db.users.update_one({"_id": user["_id"]}, {"$set": {"period.used.lesson_plans": 40}}))
        with pytest.raises(plans.PlanRefusal):
            run_db(plans.reserve(doc(backend, client.user["id"]), "lesson_plans", "hit-1"))
        hit = run_db(backend.db.db.plan_limit_hits.find_one({"user_id": client.user["id"]}))
        assert hit["kind"] == "ustaz" and hit["period_day"] == 1

    def test_the_metrics_page_counts_them(self, backend, account, plans):
        admin = _staff(backend, account)
        early, late = account(), account()
        for client, day in ((early, 3), (late, 25)):
            user = run_db(plans.grant(doc(backend, client.user["id"]), plans.USTAZ_MONTH))
            start = datetime.now(timezone.utc) - timedelta(days=day - 1)
            run_db(backend.db.db.users.update_one({"_id": user["_id"]}, {"$set": {
                "period.used.lesson_plans": 40, "period.start": start}}))
            with pytest.raises(plans.PlanRefusal):
                run_db(plans.reserve(doc(backend, client.user["id"]), "lesson_plans", f"m-{day}"))

        body = admin.get("/api/admin/metrics?days=7").json()
        row = next(r for r in body["limit_hits"] if r["resource"] == "lesson_plans")
        assert row["ustaz_users"] >= 2 and row["early_users"] >= 1
        assert row["early_users"] < row["ustaz_users"]
        assert body["early_period_day"] == 15


class TestPlanSettings:
    def test_a_stored_plan_gets_new_settings_but_keeps_its_numbers(self, backend, plans):
        old = {"_id": plans.USTAZ_MONTH, "kind": "subscription", "title": "Ұстаз · ай", "price": 4990,
               "period": "month",
               "limits": {"videos": 20, "assessments": 40, "quiz_drafts": None, "lesson_plans": 40},
               "carryover_cap": 20, "fair_use": {}, "max_effort": "high", "watermark": False,
               "docx_signature": False}
        try:
            run_db(backend.db.db.plans.delete_one({"_id": plans.USTAZ_MONTH}))
            run_db(backend.db.db.plans.insert_one(old))
            run_db(plans.seed_plans())
            stored = run_db(backend.db.db.plans.find_one({"_id": plans.USTAZ_MONTH}))
            assert stored["limits"]["assessments"] == 40  # the admin's number stays
            assert stored["carryover_caps"] == {"assessments": 20}  # a setting it never had arrives
            assert stored["fair_use"] == {}  # one an admin emptied stays empty
        finally:
            run_db(backend.db.db.plans.delete_one({"_id": plans.USTAZ_MONTH}))

    def test_carry_over_and_fair_use_are_edited_from_the_panel(self, backend, account):
        admin = _staff(backend, account)
        try:
            response = admin.put("/api/admin/plans/ustaz_month", json={
                "carryover_caps": {"assessments": 5}, "fair_use": {"quiz_drafts": None}, "reason": "Тест"})
            assert response.status_code == 200, response.text
            assert response.json()["carryover_caps"] == {"assessments": 5}
            assert response.json()["fair_use"] == {}
            refused = admin.put("/api/admin/plans/ustaz_month",
                                json={"carryover_caps": {"videos": 5}, "reason": "Тест"})
            assert refused.status_code == 400  # videos have their own field
        finally:
            run_db(backend.db.db.plans.delete_one({"_id": "ustaz_month"}))


# ------------------------------------------------------- two variants --


class TestTwoVariants:
    """1-нұсқа and 2-нұсқа: one generation, one paper off the plan."""

    @pytest.fixture(autouse=True)
    def _two(self, backend, monkeypatch):
        from app.ws.manager import agent_manager

        self.asked = []

        async def paper(spec, timeout):
            self.asked.append(spec)
            if spec["count"] == 3:  # a swap of one question
                return {"questions": [{"text": "Жаңа сұрақ", "marks": 7}]}
            out = {"questions": [{"text": "Бірінші нұсқа 1", "marks": 2},
                                 {"text": "Бірінші нұсқа 2", "marks": 3}]}
            if spec.get("variants") == 2:
                out["variant2"] = [{"text": "Екінші нұсқа 1", "marks": 2},
                                   {"text": "Екінші нұсқа 2", "marks": 3}]
            return out

        monkeypatch.setattr(agent_manager, "request_assessment", paper)

    def test_two_variants_are_one_paper_off_the_plan(self, account):
        client = account()
        response = client.post("/api/assessments", json={**SPEC, "variants": 2})
        assert response.status_code == 201, response.text
        paper = response.json()
        assert [q["text"] for q in paper["variant2"]] == ["Екінші нұсқа 1", "Екінші нұсқа 2"]
        assert all(q["id"] for q in paper["variant2"])
        assert self.asked[-1]["variants"] == 2
        assert client.get("/api/plan").json()["plan"]["used"]["assessments"] == 1

    def test_one_variant_unless_asked(self, account):
        paper = account().post("/api/assessments", json=SPEC).json()
        assert paper["variant2"] == [] and self.asked[-1]["variants"] == 1

    def test_the_second_variant_is_edited_on_its_own(self, account):
        client = account()
        paper = client.post("/api/assessments", json={**SPEC, "variants": 2}).json()
        edited = [{**q, "text": "Өзгертілген"} if i == 0 else q for i, q in enumerate(paper["variant2"])]

        saved = client.put(f"/api/assessments/{paper['id']}/questions",
                           json={"questions": edited, "variant": 2}).json()

        assert saved["variant2"][0]["text"] == "Өзгертілген"
        assert saved["questions"] == paper["questions"]

    def test_a_one_variant_paper_has_no_second_to_edit(self, account):
        client = account()
        paper = client.post("/api/assessments", json=SPEC).json()
        response = client.put(f"/api/assessments/{paper['id']}/questions",
                              json={"questions": paper["questions"], "variant": 2})
        assert response.status_code == 404

    def test_a_swap_in_the_second_variant_keeps_the_marks(self, account):
        client = account()
        paper = client.post("/api/assessments", json={**SPEC, "variants": 2}).json()
        target = paper["variant2"][1]

        swapped = client.post(
            f"/api/assessments/{paper['id']}/questions/{target['id']}/regenerate?variant=2").json()

        assert swapped["variant2"][1] == {"id": target["id"], "text": "Жаңа сұрақ", "marks": 3}
        assert swapped["questions"] == paper["questions"]

    def test_the_docx_has_both_sheets(self, account):
        from docx import Document

        client = account()
        paper = client.post("/api/assessments", json={**SPEC, "variants": 2}).json()
        content = client.get(f"/api/assessments/{paper['id']}/docx").content
        body = "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)
        assert "1-нұсқа" in body and "2-нұсқа" in body
        assert body.index("Бірінші нұсқа 2") < body.index("2-нұсқа") < body.index("Екінші нұсқа 1")
        assert body.count("Жалпы: 5 балл") == 2
