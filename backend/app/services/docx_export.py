"""An assessment as a Word document, ready to print or edit.

Built here, in the request, rather than queued for the exporter: a paper of
twenty questions is a few kilobytes and a few milliseconds, and the exporter's
jobs are built around a source video this has none of.

The text follows the paper's own language (Kazakh or Russian), not the
interface's: it is the students' sheet. A trial paper carries "Akron арқылы
жасалды" in the footer of every page; Ұстаз papers carry nothing.

A paper with two variants prints each on its own page, headed 1-нұсқа and
2-нұсқа, so the teacher prints the file and hands out alternate sheets.
"""

import io
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor

SIGNATURE = "Akron арқылы жасалды"

_TYPE = {
    "kk": {"sor": "БЖБ", "soch": "ТЖБ", "quiz": "Тест"},
    "ru": {"sor": "СОР", "soch": "СОЧ", "quiz": "Тест"},
}
_SUBJECT = {
    "kk": {"math": "Математика", "geometry": "Геометрия", "physics": "Физика",
           "chemistry": "Химия", "biology": "Биология", "informatics": "Информатика"},
    "ru": {"math": "Математика", "geometry": "Геометрия", "physics": "Физика",
           "chemistry": "Химия", "biology": "Биология", "informatics": "Информатика"},
}
_TEXT = {
    "kk": {
        "grade": "{n}-сынып",
        "topic": "Тақырып: {t}",
        "student": "Оқушының аты-жөні: ______________________   Сынып: ________   Күні: ________",
        "marks": "({n} балл)",
        "total": "Жалпы: {n} балл",
        "variant": "{n}-нұсқа",
    },
    "ru": {
        "grade": "{n} класс",
        "topic": "Тема: {t}",
        "student": "ФИО ученика: ______________________   Класс: ________   Дата: ________",
        "marks": "({n} б.)",
        "total": "Всего: {n} баллов",
        "variant": "Вариант {n}",
    },
}


def _lang(paper: Dict[str, Any]) -> str:
    return paper.get("lang") if paper.get("lang") in _TEXT else "kk"


def title_of(paper: Dict[str, Any]) -> str:
    lang = _lang(paper)
    kind = _TYPE[lang].get(paper.get("type", ""), "")
    subject = _SUBJECT[lang].get(paper.get("subject", ""), paper.get("subject", ""))
    grade = _TEXT[lang]["grade"].format(n=paper.get("grade", ""))
    return " · ".join(part for part in (kind, subject, grade) if part)


def assessment_docx(paper: Dict[str, Any], *, signature: bool) -> bytes:
    lang = _lang(paper)
    document = Document()

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    variants = [paper.get("questions") or []]
    if paper.get("variant2"):
        variants.append(paper["variant2"])
    for number, questions in enumerate(variants, start=1):
        if number > 1:
            document.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        _sheet(document, paper, questions, lang, number if len(variants) > 1 else None)

    if signature:
        for section in document.sections:
            footer = section.footer.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            mark = footer.add_run(SIGNATURE)
            mark.font.size = Pt(9)
            mark.font.color.rgb = RGBColor(0x80, 0x88, 0x99)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _sheet(document: Any, paper: Dict[str, Any], questions: List[Dict[str, Any]], lang: str,
           variant: Optional[int]) -> None:
    """One student's sheet: the heading, the name line, the questions, the total."""
    text = _TEXT[lang]
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title_of(paper))
    run.bold = True
    run.font.size = Pt(14)

    if variant:
        label = document.add_paragraph()
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label.add_run(text["variant"].format(n=variant)).bold = True

    if paper.get("topic"):
        topic = document.add_paragraph(text["topic"].format(t=paper["topic"]))
        topic.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph(text["student"])

    total = 0
    for number, question in enumerate(questions, start=1):
        marks = int(question.get("marks") or 0)
        total += marks
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.add_run(f"{number}. ").bold = True
        paragraph.add_run(str(question.get("text") or "").strip())
        if marks:
            paragraph.add_run(" " + text["marks"].format(n=marks)).italic = True

    closing = document.add_paragraph()
    closing.paragraph_format.space_before = Pt(16)
    closing.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    closing.add_run(text["total"].format(n=total)).bold = True
